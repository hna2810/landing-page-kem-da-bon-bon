import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Sheet 1: Danh sách video
ws1 = wb.active
ws1.title = 'Top 20 Video Viral Mẹ Bé'

# Freeze panes at row 2 so headers stay fixed when scrolling
ws1.freeze_panes = 'A2'

# Headers
headers = [
    'STT', 
    'Nhóm Chủ Đề', 
    'Link Video Douyin', 
    'Ngày Đăng', 
    'Tim Douyin (Ước tính)', 
    'View Ước Tính', 
    'Tóm Tắt Kịch Bản & Điểm Nổi Bật', 
    'Ý Tưởng Ứng Dụng Làm Ads', 
    'Sản Phẩm Phù Hợp'
]

ws1.append(headers)

data = [
    # Nhóm 1: AI Video > 5M Views
    (1, '🤖 AI Mẹ & Bé (> 5M View)', 'https://www.douyin.com/video/7695169182837469386', '10/10/2026', '30w+ tim (300k)', '> 10.000.000', 
     'AI Mengwa (AI萌娃) nhảy múa & tương tác: Bé 3D siêu thực nhảy nhót khớp nhạc, visual siêu cute đốn tim.', 
     'Hook 3s đầu: Bé nhảy vui vẻ sau khi uống sữa hoặc mang bỉm êm ái, kéo tỷ lệ giữ chân người xem cực cao.', 
     'Sữa công thức, tã bỉm, bình sữa cao cấp'),
     
    (2, '🤖 AI Mẹ & Bé (> 5M View)', 'https://www.douyin.com/video/7688611324369134976', '23/09/2026', '25w+ tim (250k)', '> 8.000.000', 
     'AI Baby đối thoại như người lớn: Bé sơ sinh biểu cảm bà cụ non trò chuyện với mẹ, tình huống bất ngờ gây cười.', 
     'Bé "than thở" mẹ pha sữa lâu quá khiến bé đói ➡️ Dẫn dắt sang máy pha sữa tự động 5s.', 
     'Máy pha sữa, bình đun giữ nhiệt, nôi điện'),
     
    (3, '🤖 AI Mẹ & Bé (> 5M View)', 'https://www.douyin.com/video/7686124139061371849', '16/09/2026', '20w+ tim (200k)', '> 7.500.000', 
     'Kling AI - Bé sơ sinh biểu cảm siêu thực: Từng sợi tóc và làn da em bé chân thực như quay máy quay rạp phim.', 
     'Cận cảnh làn da em bé mỏng manh ➡️ Quảng cáo sản phẩm an toàn tuyệt đối cho trẻ sơ sinh.', 
     'Tã bỉm organic, khăn sữa sơ sinh, kem dưỡng da bé'),
     
    (4, '🤖 AI Mẹ & Bé (> 5M View)', 'https://www.douyin.com/video/7686051037417599598', '16/09/2026', '18w+ tim (180k)', '> 6.000.000', 
     'Kling AI - Bé sơ sinh "đòi quà" bố mẹ: Em bé mở miệng kể danh sách quà muốn bố mẹ tặng trước khi chào đời.', 
     'Bé đọc danh sách đồ bố mẹ "bắt buộc phải sắm" ➡️ Kịch bản bán combo trọn gói mẹ bé.', 
     'Combo đón bé chào đời, giỏ đồ đi sinh, đồ chơi sơ sinh'),
     
    (5, '🤖 AI Mẹ & Bé (> 5M View)', 'https://www.douyin.com/video/7685605652786292404', '15/09/2026', '15w+ tim (150k)', '> 5.500.000', 
     'Jimeng AI - Thiên thần nhí chúc phúc: Hình ảnh em bé thiên thần lung linh, âm nhạc êm dịu cảm động.', 
     'Đánh vào tình mẫu tử thiêng liêng và sự mong đợi của mẹ bầu ➡️ Kích thích đăng ký dịch vụ cao cấp.', 
     'Gói sinh trọn gói, bảo hiểm mẹ bé, spa chăm sóc bầu'),

    # Nhóm 2: ASMR Chăm con & Nuôi con nhàn tênh
    (6, '🍼 ASMR Chăm Con', 'https://www.douyin.com/video/7690249939255618152', '27/09/2026', '10w+ tim (100k)', '> 3.200.000', 
     'Vlog một ngày chăm con ASMR không lời: Thu âm tiếng bóc bỉm, tiếng rót nước, tiếng máy vận hành êm ái.', 
     'Nuôi con nhàn tênh nhờ công nghệ ➡️ Khoe trọn tính năng máy móc vận hành ban đêm không làm bé tỉnh giấc.', 
     'Máy tiệt trùng sấy UV, bình đun giữ nhiệt, nôi rung'),
     
    (7, '🍼 ASMR Chăm Con', 'https://www.douyin.com/video/7689443262788626825', '25/09/2026', '8w+ tim (80k)', '> 2.800.000', 
     'Buổi sáng chăm bé sơ sinh: Cận cảnh các bước vệ sinh mắt mũi, thay bỉm, mát-xa bụng chống đầy hơi.', 
     'Quy trình chăm sóc chuẩn y khoa ➡️ Giới thiệu sản phẩm theo từng bước buổi sáng tạo độ tin cậy.', 
     'Nước muối tép sinh lý, dầu massage bé, tã dán organic'),
     
    (8, '🍼 ASMR Chăm Con', 'https://www.douyin.com/video/7686434900283932650', '17/09/2026', '6w+ tim (60k)', '> 1.900.000', 
     'Góc nhìn thứ nhất (POV) mẹ chăm con: Camera gắn trước ngực quay thao tác 1 tay bế con, 1 tay lấy đồ.', 
     'Nhấn mạnh vào tính năng thiết kế công thái học, giúp mẹ bỉm thao tác rảnh tay dễ dàng.', 
     'Bình sữa chống sặc cổ nghiêng, gối chống trào ngược'),
     
    (9, '🍼 ASMR Chăm Con', 'https://www.douyin.com/video/7692361343018909794', '03/10/2026', '7w+ tim (70k)', '> 2.100.000', 
     'Chăm bé một mình không stress: Mẹ bỉm hiện đại setup không gian phòng ngủ và góc thay đồ siêu tiện lợi.', 
     'Góc phòng bỉm sữa ngăn nắp giúp mẹ tiết kiệm 50% thời gian tìm đồ.', 
     'Kệ bỉm sữa đa năng 3 tầng có bánh xe, giỏ gấp gọn'),
     
    (10, '🍼 ASMR Chăm Con', 'https://www.douyin.com/video/7686053667733144235', '16/09/2026', '5w+ tim (50k)', '> 1.500.000', 
     'Nhật ký chăm sóc bé ban đêm: Giải quyết nỗi sợ thức đêm của mẹ bỉm nhờ các thiết bị chuyên dụng.', 
     'Nỗi đau thức đêm dậy pha sữa mò mẫm trong tối ➡️ Giải pháp máy pha sữa 1 chạm không cần bật đèn sáng.', 
     'Đèn ngủ cảm ứng chống chói mắt, máy pha sữa đêm 1 chạm'),

    # Nhóm 3: Ăn dặm & Dụng cụ chế biến
    (11, '🥣 Nấu Ăn Dặm', 'https://www.douyin.com/video/7691969512995835162', '02/10/2026', '8w+ tim (80k)', '> 2.500.000', 
     'Món ăn dặm bắt mắt kích thích vị giác: Làm bánh flan rau củ, cháo tôm bí đỏ màu sắc rực rỡ, bé ăn thun thút.', 
     'Cách làm món ngon giúp bé hết biếng ăn ➡️ Giới thiệu dụng cụ chế biến mini tiện lợi.', 
     'Máy xay ăn dặm mini 8 lưỡi, khuôn bánh silicone'),
     
    (12, '🥣 Nấu Ăn Dặm', 'https://www.douyin.com/video/7690844513597754634', '29/09/2026', '5.5w+ tim (55k)', '> 1.800.000', 
     'Chế biến nhanh 10 phút cho mẹ bận rộn: Công thức cấp đông nước dùng dashi và cháo hạt cho bé.', 
     'Giải pháp chuẩn bị đồ ăn cả tuần cho mẹ bỉm đi làm nhưng con vẫn ăn đồ tươi ngon.', 
     'Hộp trữ đông chia vạch, nồi nấu cháo cách thủy hẹn giờ'),

    # Nhóm 4: Giỏ đồ đi sinh & Combo sơ sinh
    (13, '🏥 Giỏ Đồ Đi Sinh', 'https://www.douyin.com/video/7689894222425492763', '26/09/2026', '10w+ tim (100k)', '> 3.000.000', 
     'Bóc tách vali đi sinh (Cần vs Thừa): Chỉ ra đồ viện phát và những món thực sự mẹ cần mang theo.', 
     'Cứu cánh ví tiền mẹ bầu, không lo mua phải đồ vô dụng ➡️ Giới thiệu giỏ đồ đi sinh tối giản.', 
     'Combo giỏ đồ đi sinh chia túi zip, tã dán, khăn khô'),
     
    (14, '🏥 Giỏ Đồ Đi Sinh', 'https://www.douyin.com/video/7687180844741504506', '19/09/2026', '5.5w+ tim (55k)', '> 1.700.000', 
     'Chồng chuẩn bị giỏ đồ đi sinh cho vợ: Ông bố dán nhãn từng túi đồ để vào viện không bị cuống.', 
     'Đánh trúng tâm lý người chồng chu đáo; thúc đẩy các cặp đôi mua trọn set có phân loại sẵn.', 
     'Set đồ sơ sinh đóng gói theo set, đồ lót giấy cho mẹ'),

    # Nhóm 5: Xe đẩy & Đồ ra ngoài
    (15, '🚼 Dụng Cụ Ra Ngoài', 'https://www.douyin.com/video/7694274948583935354', '08/10/2026', '7.5w+ tim (75k)', '> 2.400.000', 
     'Gấp gọn xe đẩy 1 giây bằng một tay: Tay bế con, chân gạt nhẹ tay cầm bấm nút gấp xe như vali.', 
     'Show tính năng trực quan giải quyết cảnh mẹ đi chơi một mình vừa ôm con vừa vác xe cồng kềnh.', 
     'Xe đẩy du lịch gấp gọn dưới 5kg, vali kéo cho bé'),

    # Nhóm 6: Hook em bé triệu view (3 giây đầu)
    (16, '👶 Hook Em Bé Triệu View', 'https://www.douyin.com/video/7694703932287675903', '09/10/2026', '15w+ tim (150k)', '> 4.500.000', 
     'Bé làm nũng ánh mắt long lanh đòi bế: Biểu cảm dễ thương tuyệt đối giữ chân người xem.', 
     'Chuyển cảnh: "Bé bám mẹ cả ngày thế này thì làm sao làm việc? ➡️ Địu trợ lực chống mỏi lưng".', 
     'Địu trợ lực công thái học, nôi rung tự động'),
     
    (17, '👶 Hook Em Bé Triệu View', 'https://www.douyin.com/video/7694322017751368043', '08/10/2026', '12w+ tim (120k)', '> 3.800.000', 
     'Bé nghịch bôi bẩn quần áo và sàn nhà: Tình huống dở khóc dở cười khiến bố mẹ gật gù đồng cảm.', 
     'Chuyển cảnh: "Đừng cấm con tự do nghịch ngợm ➡️ Khăn ướt 0% cồn & nước giặt enzyme đánh bay vết ố".', 
     'Khăn ướt an toàn, nước giặt enzyme cho bé'),
     
    (18, '👶 Hook Em Bé Triệu View', 'https://www.douyin.com/video/7692616429375820262', '04/10/2026', '10w+ tim (100k)', '> 3.100.000', 
     'Bé ngủ gật ngay trên bàn ăn dặm: Đầu gật gù miệng vẫn ngậm thìa thức ăn, cực kỳ hài hước.', 
     'Chuyển cảnh: "Bé ăn ngon miệng là ngủ say ➡️ Ghế ăn dặm ngả 3 nấc biến thành giường ngủ".', 
     'Ghế ăn dặm ngả nằm, yếm máng silicone'),
     
    (19, '👶 Hook Em Bé Triệu View', 'https://www.douyin.com/video/7691270236385215808', '30/09/2026', '8w+ tim (80k)', '> 2.600.000', 
     'Bé nhăn nhó khi lần đầu nếm vị chanh: Biểu cảm ngạc nhiên và khó đỡ gây sốt cộng đồng mạng.', 
     'Chuyển cảnh: "Làm sao giúp con làm quen đồ ăn mới không quấy khóc? ➡️ Set thìa núm ăn dặm chống hóc".', 
     'Túi nhai ăn dặm chống hóc, set thìa mềm silicon'),
     
    (20, '👶 Hook Em Bé Triệu View', 'https://www.douyin.com/video/7691959333383200936', '02/10/2026', '7w+ tim (70k)', '> 2.200.000', 
     'Bố bỉm lóng ngóng thay tã làm rơi đồ: Hình ảnh ông bố vụng về nhưng rất yêu con tạo nhiều bình luận.', 
     'Chuyển cảnh: "Giao con cho bố không lo nếu có dụng cụ này ➡️ Bàn thay tã chống mỏi lưng".', 
     'Bàn thay tã chống mỏi lưng di động, giỏ tã bỉm')
]

for row in data:
    ws1.append(list(row))

# Formatting Sheet 1
header_fill = PatternFill(start_color='E11D48', end_color='E11D48', fill_type='solid') # Rose Red
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
cell_font = Font(name='Segoe UI', size=10)
link_font = Font(name='Segoe UI', size=10, color='2563EB', underline='single')
bold_font = Font(name='Segoe UI', size=10, bold=True)

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

ai_fill = PatternFill(start_color='F5F3FF', end_color='F5F3FF', fill_type='solid') # Light Purple
asmr_fill = PatternFill(start_color='FFF1F2', end_color='FFF1F2', fill_type='solid') # Light Rose
wean_fill = PatternFill(start_color='FEFCE8', end_color='FEFCE8', fill_type='solid') # Light Yellow
hook_fill = PatternFill(start_color='F0FDF4', end_color='F0FDF4', fill_type='solid') # Light Green

for col_idx, cell in enumerate(ws1[1], start=1):
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws1.row_dimensions[1].height = 28

for row_idx in range(2, len(data) + 2):
    ws1.row_dimensions[row_idx].height = 45
    cat_text = ws1.cell(row=row_idx, column=2).value
    
    current_fill = None
    if 'AI' in cat_text:
        current_fill = ai_fill
    elif 'ASMR' in cat_text:
        current_fill = asmr_fill
    elif 'Nấu Ăn Dặm' in cat_text:
        current_fill = wean_fill
    elif 'Hook' in cat_text:
        current_fill = hook_fill

    for col_idx in range(1, len(headers) + 1):
        cell = ws1.cell(row=row_idx, column=col_idx)
        cell.border = thin_border
        cell.font = cell_font
        if current_fill:
            cell.fill = current_fill
            
        if col_idx in [1, 4]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif col_idx in [5, 6]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = bold_font
        elif col_idx == 3: # Link
            raw_url = cell.value
            cell.hyperlink = raw_url
            cell.value = 'Mở Video Douyin'
            cell.font = link_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

# Set Column Widths
col_widths = {
    1: 6,   # STT
    2: 24,  # Nhóm Chủ Đề
    3: 18,  # Link Video
    4: 13,  # Ngày Đăng
    5: 22,  # Tim Douyin
    6: 18,  # View Ước Tính
    7: 42,  # Tóm Tắt
    8: 42,  # Ý Tưởng Ads
    9: 30   # Sản Phẩm
}
for col_idx, width in col_widths.items():
    ws1.column_dimensions[get_column_letter(col_idx)].width = width

# Sheet 2: Kịch Bản Ads Chuẩn
ws2 = wb.create_sheet(title='4 Công Thức Kịch Bản Ads')
ws2.freeze_panes = 'A2'
ws2.append(['Mã Kịch Bản', 'Tên Mô Tuýp Kịch Bản', 'Cấu Trúc Thời Lượng (30-45s)', 'Tâm Lý Đánh Trúng', 'Sản Phẩm Tối Ưu'])

kb_data = [
    ('KB-01', 'ASMR / Chăm Con Nhàn Tênh (Immersive Routine)', 
     '0-3s: Âm thanh giật gân (tiếng bóc vỏ, tiếng rót nước, tiếng máy chạy êm).\n4-15s: Thao tác nhẹ nhàng, góc quay cận cảnh đồ dùng công nghệ.\n16-30s: Bé nằm ngoan/uống sữa ngon, mẹ thư thái uống trà.\nCTA: Để lại bình luận nhận link máy.', 
     'Đánh vào sự kiệt sức của mẹ bỉm sau sinh; khát khao được nghỉ ngơi và có người đỡ đần.', 
     'Máy pha sữa, máy tiệt trùng UV, nôi điện ru ngủ, máy hút sữa rảnh tay'),
     
    ('KB-02', 'Nỗi Đau ➡️ Giải Pháp (Pain Point to Solution)', 
     '0-3s: Cảnh mẹ bù đầu tóc rối nửa đêm dậy pha sữa hoặc bé quấy khóc dữ dội.\n4-10s: Đặt câu hỏi: "Sao người ta nuôi con nhàn thế còn mình thì phát điên?".\n11-25s: Tiết lộ vũ khí bí mật (Sản phẩm quảng cáo).\nCTA: Bấm giỏ hàng ngay hôm nay để nhận ưu đãi cho mẹ bỉm.', 
     'Đánh vào cảm giác tự ti, áp lực làm mẹ, nỗi sợ con không được chăm sóc tốt nhất.', 
     'Bình sữa chống sặc, gối chống trào ngược, bàn thay tã chống mỏi lưng'),
     
    ('KB-03', 'Tránh Bẫy Mua Sắm (Bóc Phốt Đồ Lãng Phí)', 
     '0-3s: "3 món đồ mẹ bầu nào cũng mua nhưng 100% vứt xó lãng phí tiền!".\n4-15s: Kể tên 2 món vô dụng ai cũng dính (để tạo uy tín khách quan).\n16-30s: Giới thiệu món thứ 3: "Thay vào đó, chỉ cần sắm đúng món này là dùng từ sơ sinh tới 3 tuổi".\nCTA: Giảm ngay 20% cho các mẹ chuẩn bị đi sinh.', 
     'Đánh vào tâm lý sợ mất tiền oan, mẹ bầu chưa có kinh nghiệm dễ bị lúng túng.', 
     'Combo giỏ đồ đi sinh tối giản, xe đẩy gấp gọn, ghế ăn dặm đa năng'),
     
    ('KB-04', 'AI Baby / Em Bé Tương Tác Hài Hước', 
     '0-3s: Bé sơ sinh biểu cảm bà cụ non / AI nói tiếng người giữ chân tuyệt đối.\n4-15s: Bé "tuyên bố" những tiêu chuẩn khắt khe về đồ ăn/bỉm sữa con muốn mẹ dùng.\n16-30s: Mẹ đưa đúng sản phẩm, bé lập tức cười tươi / nhảy múa vui nhộn.\nCTA: Mẹ thông thái chọn ngay cho con.', 
     'Tạo tiếng cười, giảm bớt tính thương mại gượng gạo của quảng cáo truyền thống.', 
     'Tã bỉm trẻ em, sữa dinh dưỡng, siro ăn ngon, kem bôi da trẻ em')
]

for row in kb_data:
    ws2.append(list(row))

for col_idx, cell in enumerate(ws2[1], start=1):
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws2.row_dimensions[1].height = 28

for row_idx in range(2, len(kb_data) + 2):
    ws2.row_dimensions[row_idx].height = 65
    for col_idx in range(1, 6):
        cell = ws2.cell(row=row_idx, column=col_idx)
        cell.border = thin_border
        cell.font = cell_font
        if col_idx in [1, 2]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = bold_font
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

kb_widths = {1: 15, 2: 30, 3: 50, 4: 35, 5: 35}
for col_idx, width in kb_widths.items():
    ws2.column_dimensions[get_column_letter(col_idx)].width = width

# Sheet 3: Hướng Dẫn & Spy Công Cụ
ws3 = wb.create_sheet(title='Từ Khóa & Công Cụ Tải Video')
ws3.freeze_panes = 'A2'
ws3.append(['STT', 'Từ Khóa Tiếng Trung', 'Phiên Âm Pinyin', 'Dịch Nghĩa Tiếng Việt', 'Mục Đích Dùng Trên Douyin/Xiaohongshu'])

keywords = [
    (1, '沉浸式带娃', 'Chénjìn shì dài wá', 'Chăm con kiểu ASMR / đắm chìm', 'Tìm video đồ gia dụng hiện đại, máy móc, phòng ngủ sang xịn'),
    (2, '母婴好物', 'Mǔyīng hǎowù', 'Đồ tốt mẹ và bé', 'Tìm video review thực tế, unboxing sản phẩm mới'),
    (3, '带娃神器', 'Dài wá shénqì', 'Thần khí nuôi con', 'Tìm đồ thông minh rảnh tay cho mẹ (xe đẩy, nôi, ghế ăn, địu)'),
    (4, '避坑指南', 'Bì kēng zhǐnán', 'Hướng dẫn tránh bẫy mua sắm', 'Tìm kịch bản bóc phốt đồ vô dụng vs khen đồ chân ái'),
    (5, '待产包清单', 'Dàichǎn bāo qīngdān', 'Giỏ đồ đi sinh bệnh viện', 'Tìm ý tưởng đóng gói combo sơ sinh bán trọn gói'),
    (6, '辅食打卡 / 辅食教程', 'Fǔshí dǎkǎ / jiàochéng', 'Hướng dẫn nấu đồ ăn dặm', 'Tìm video làm bánh, cháo bắt mắt bán máy xay và đồ ăn dặm'),
    (7, '人类幼崽', 'Rénlèi yòuzǎi', 'Ấu tể loài người (em bé cute)', 'Tìm tình huống hài hước, bé làm nũng làm Hook 3s đầu'),
    (8, 'AI萌娃 / 即梦AI', 'AI méngwá / Jímèng AI', 'Em bé AI ảo triệu view', 'Tìm nguồn cảm hứng làm video AI giữ chân người xem')
]

for row in keywords:
    ws3.append(list(row))

for col_idx, cell in enumerate(ws3[1], start=1):
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 28

for row_idx in range(2, len(keywords) + 2):
    ws3.row_dimensions[row_idx].height = 30
    for col_idx in range(1, 6):
        cell = ws3.cell(row=row_idx, column=col_idx)
        cell.border = thin_border
        cell.font = cell_font
        if col_idx in [1, 3]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif col_idx == 2:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = bold_font
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

kw_widths = {1: 8, 2: 24, 3: 24, 4: 30, 5: 45}
for col_idx, width in kw_widths.items():
    ws3.column_dimensions[get_column_letter(col_idx)].width = width

# Add notes on web tools to Sheet 3
ws3.cell(row=12, column=1, value='💡 CÔNG CỤ TẢI VIDEO DOUYIN KHÔNG DÍNH WATERMARK (LOGO):').font = bold_font
tools = [
    ('1. Trang web Tikwm:', 'https://www.tikwm.com (Dán link Douyin vào tải mp4 HD sạch logo)'),
    ('2. Trang web SnapTik:', 'https://snaptik.app (Hỗ trợ tải video Douyin/Tiktok không watermark)'),
    ('3. Bot Telegram tiện lợi:', 'Tìm kiếm bot @DouyinSaveBot hoặc @SaveRestrictedBot gửi link tải trực tiếp')
]
for r_idx, (t_name, t_url) in enumerate(tools, start=13):
    ws3.cell(row=r_idx, column=2, value=t_name).font = bold_font
    ws3.cell(row=r_idx, column=3, value=t_url)

file_path = r'c:\Users\Admin\Desktop\New folder (2)\Tong_Hop_Video_Me_Be_Viral_30_Ngay.xlsx'
wb.save(file_path)
print('SUCCESS: Created Excel file at:', file_path)
